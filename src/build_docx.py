# -*- coding: utf-8 -*-
"""Build How-Tech-Moves-10-Video-Concepts.docx from the shared content model."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import content

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(
    HERE, "..", "How-Tech-Moves-10-Video-Concepts.docx"))

# ---------------------------------------------------------------- design tokens
SERIF = "Georgia"          # headings
SANS = "Calibri"           # body
MONO = "Consolas"          # code-style labels block

BODY_C = RGBColor(0x1F, 0x1E, 0x1D)     # near-black, never pure black
ACCENT = RGBColor(0xC0, 0x56, 0x2F)     # terracotta clay (single accent)
META_C = RGBColor(0x6B, 0x6A, 0x66)     # warm grey
HAIRLINE = "DEDCD1"
CODE_BG = "F0EEE6"
CODE_EDGE = "E8E6DC"

# CT_PPrBase child order (schema) so injected pBdr/shd land legally
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
]


def insert_ordered(pPr, el):
    """Insert el into pPr respecting the CT_PPrBase sequence."""
    tag = el.tag.split("}")[1]
    idx = PPR_ORDER.index(tag)
    for child in pPr:
        ctag = child.tag.split("}")[1]
        if ctag in PPR_ORDER and PPR_ORDER.index(ctag) > idx:
            child.addprevious(el)
            return
    pPr.append(el)


def set_run_font(run, name, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    # letter-spacing normal (explicit, no tracking)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), "0")
    rPr.append(sp)


def set_borders(p, sides=("top", "left", "bottom", "right"), color=HAIRLINE,
                sz="4", space=6):
    pPr = p._p.get_or_add_pPr()
    for el in pPr.findall(qn("w:pBdr")):
        pPr.remove(el)
    pBdr = OxmlElement("w:pBdr")
    for side in sides:
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), str(space))
        e.set(qn("w:color"), color)
        pBdr.append(e)
    insert_ordered(pPr, pBdr)


def set_shading(p, fill):
    pPr = p._p.get_or_add_pPr()
    for el in pPr.findall(qn("w:shd")):
        pPr.remove(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(pPr, shd)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run()
    set_run_font(run, SANS, 8.5, META_C)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._element.append(fld_begin)
    run._element.append(instr)
    run._element.append(fld_end)


# ---------------------------------------------------------------- build
doc = Document()

# A4, 2.2 cm margins
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)
sec.footer_distance = Cm(1.2)
add_page_number_footer(sec)

# Normal = body
normal = doc.styles["Normal"]
normal.font.name = SANS
normal.font.size = Pt(11)
normal.font.color.rgb = BODY_C
normal.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
normal.element.rPr.rFonts.set(qn("w:cs"), SANS)
pf = normal.paragraph_format
pf.line_spacing = 1.3
pf.space_after = Pt(8)
pf.space_before = Pt(0)


def mkstyle(name, font, size, color, bold=False, before=0, after=8,
            line=1.3, keep_next=False, keep_lines=False):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = bold
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.element.rPr.rFonts.set(qn("w:cs"), font)
    f = st.paragraph_format
    f.line_spacing = line
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.keep_with_next = keep_next
    f.keep_together = keep_lines
    return st


S_H1 = mkstyle("HTM H1", SERIF, 26, ACCENT, False, 0, 6, 1.15, False, True)
S_SUB = mkstyle("HTM Subtitle", SANS, 14, META_C, False, 6, 4, 1.35, False, True)
S_H2 = mkstyle("HTM H2", SERIF, 15, ACCENT, True, 21, 8, 1.2, False, True)
S_H3 = mkstyle("HTM H3", SERIF, 11.5, ACCENT, True, 12, 2, 1.25, False, True)
S_BODY = mkstyle("HTM Body", SANS, 11, BODY_C, False, 0, 8, 1.3, False, True)
S_LEAD = mkstyle("HTM Lead", SANS, 11, BODY_C, True, 8, 4, 1.3, True, True)
S_CODE = mkstyle("HTM Code", MONO, 9.5, BODY_C, False, 8, 12, 1.25, False, True)

# ---------------------------------------------------------------- walk blocks
group = []          # paragraph objects of the current section group
groups = []
pending_break = False

for bi, (kind, payload) in enumerate(content.BLOCKS):

    if kind == "pagebreak":
        pending_break = True
        continue

    if kind == "h1":
        p = doc.add_paragraph(style=S_H1)
        set_run_font(p.add_run(payload), SERIF, 26, ACCENT, False)
        if pending_break:
            p.paragraph_format.page_break_before = True
            pending_break = False

    elif kind == "subtitle":
        p = doc.add_paragraph(style=S_SUB)
        set_run_font(p.add_run(payload), SANS, 14, META_C, False)

    elif kind == "rule":
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1
        set_borders(p, sides=("bottom",), color=HAIRLINE, sz="4", space=0)

    elif kind == "h2":
        p = doc.add_paragraph(style=S_H2)
        set_run_font(p.add_run(payload), SERIF, 15, ACCENT, True)
        if pending_break:
            p.paragraph_format.page_break_before = True
            pending_break = False
        if group:
            groups.append(group)
        group = [p]

    elif kind == "h3":
        p = doc.add_paragraph(style=S_H3)
        set_run_font(p.add_run(payload), SERIF, 11.5, ACCENT, True)
        group.append(p)

    elif kind == "body":
        p = doc.add_paragraph(style=S_BODY)
        set_run_font(p.add_run(payload), SANS, 11, BODY_C, False)
        group.append(p)

    elif kind == "lead":
        p = doc.add_paragraph(style=S_LEAD)
        set_run_font(p.add_run(payload), SANS, 11, BODY_C, True)
        group.append(p)

    elif kind in ("olist", "bullets"):
        style = "List Number" if kind == "olist" else "List Bullet"
        for lead, text in payload:
            p = doc.add_paragraph(style=style)
            f = p.paragraph_format
            f.left_indent = Cm(0.85)
            f.first_line_indent = Cm(-0.45)
            f.space_after = Pt(4)
            f.line_spacing = 1.28
            f.keep_together = True
            if lead:
                set_run_font(p.add_run(lead), SANS, 11, BODY_C, True)
                set_run_font(p.add_run(" " + text), SANS, 11, BODY_C, False)
            else:
                set_run_font(p.add_run(text), SANS, 11, BODY_C, False)
            group.append(p)

    elif kind == "code":
        p = doc.add_paragraph(style=S_CODE)
        f = p.paragraph_format
        f.left_indent = Cm(0.15)
        f.right_indent = Cm(0.15)
        f.space_before = Pt(10)
        f.space_after = Pt(14)
        set_run_font(p.add_run(payload), MONO, 9.5, BODY_C, False)
        set_shading(p, CODE_BG)
        set_borders(p, color=CODE_EDGE, sz="4", space=7)
        group.append(p)

if group:
    groups.append(group)

# ------------------------------------------------- keep each concept together
for g in groups:
    for i, p in enumerate(g):
        p.paragraph_format.keep_together = True
        if i < len(g) - 1:
            p.paragraph_format.keep_with_next = True

# A paragraph sitting directly in front of a hard page break must not carry
# keep-with-next, or Word would drag it onto the following page and leave the
# previous page blank (this is the "THE 10 CONCEPTS" heading case).
all_paras = doc.paragraphs
for i, p in enumerate(all_paras[:-1]):
    if all_paras[i + 1].paragraph_format.page_break_before:
        p.paragraph_format.keep_with_next = False

doc.save(OUT)
print("wrote", OUT, os.path.getsize(OUT), "bytes")
print("paragraph count:", len(doc.paragraphs))
